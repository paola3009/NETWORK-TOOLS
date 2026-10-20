# -*- coding: utf-8 -*-
"""
Módulo de Exportação Excel para Assessment Big Picture.
Gera a planilha de Capacidade e Inventário de Switches com as colunas:
  - Switch (Hostname)
  - Modelo
  - Total de módulos ou de switches no stack
  - Oferta total de portas (ex: C9300-48U x 3 = 144 portas)
  - Total de APs (checados no lldp e/ou cdp)
"""

import os
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[4]
DOWNLOADS_DIR = BASE_DIR / "core" / "static" / "downloads"


def gerar_planilha_inventario_switches(
    dados_switches: List[Dict[str, Any]],
    links: List[Dict[str, Any]] = None,
    nome_arquivo: str = None
) -> str:
    """
    Gera arquivo Excel (.xlsx) contendo o inventário de switches e capacidade.
    Inclui aba com a tabela de enlaces e conexões mapeadas.
    """
    DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
    if not nome_arquivo:
        data_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        nome_arquivo = f"Assessment_Switches_Capacidade_{data_str}.xlsx"

    caminho_final = str(DOWNLOADS_DIR / nome_arquivo)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario Switches"
    ws.views.sheetView[0].showGridLines = True

    # Estilos Visuais Corporativos
    font_titulo = Font(name="Calibri", size=15, bold=True, color="0F172A")
    font_subtitulo = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="1E293B")
    font_total = Font(name="Calibri", size=11, bold=True, color="0F172A")

    fill_header = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")  # Azul Corporativo
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")     # Cinza Suave Alternado
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")   # Totalizador

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(
        top=Side(border_style="thin", color="94A3B8"),
        bottom=Side(border_style="double", color="0F172A")
    )

    # 1. Título do Relatório
    ws.merge_cells("A1:G1")
    cell_title = ws["A1"]
    cell_title.value = "INVENTÁRIO E CAPACIDADE DE SWITCHES — ASSESSMENT BIG PICTURE"
    cell_title.font = font_titulo
    cell_title.alignment = align_left
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:G2")
    cell_sub = ws["A2"]
    data_geracao = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")
    cell_sub.value = f"Gerado automaticamente em: {data_geracao} | Coleta via CDP/LLDP & Stack Assessment"
    cell_sub.font = font_subtitulo
    cell_sub.alignment = align_left
    ws.row_dimensions[2].height = 18

    # Espaço
    ws.row_dimensions[3].height = 10

    # 2. Cabeçalho das Colunas
    headers = [
        "Switch",
        "Modelo",
        "Total de módulos ou de switches no stack",
        "Oferta total de portas",
        "Total de APs (checados no lldp e/ou cdp)",
        "IP de Gerência",
        "Função na Rede"
    ]

    row_header = 4
    ws.row_dimensions[row_header].height = 28

    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=row_header, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell

    # 3. Linhas de Dados
    current_row = 5
    total_portas_geral = 0
    total_aps_geral = 0

    for idx, sw in enumerate(dados_switches):
        hostname = sw.get("hostname", f"SWITCH_{idx+1}")
        modelo = sw.get("modelo", "Desconhecido")
        stack_count = sw.get("stack_count", 1)
        portas = sw.get("total_portas", 0)
        aps = sw.get("total_aps", 0)
        ip = sw.get("ip", "")
        role = sw.get("role", "Switch")

        total_portas_geral += portas
        total_aps_geral += aps

        is_alt = (idx % 2 == 1)

        row_data = [
            (hostname, align_left),
            (modelo, align_left),
            (stack_count, align_center),
            (portas, align_center),
            (aps, align_center),
            (ip, align_center),
            (role, align_center)
        ]

        ws.row_dimensions[current_row].height = 20

        for col_idx, (val, alignment) in enumerate(row_data, 1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = font_data
            c.alignment = alignment
            c.border = border_cell
            if is_alt:
                c.fill = fill_alt

        current_row += 1

    # 4. Linha de Totais
    ws.row_dimensions[current_row].height = 24
    c_tot_label = ws.cell(row=current_row, column=1, value="TOTAL GERAL")
    c_tot_label.font = font_total
    c_tot_label.alignment = align_left
    c_tot_label.fill = fill_total
    c_tot_label.border = border_total

    # Coluna 2 (Modelo)
    c_sw_count = ws.cell(row=current_row, column=2, value=f"{len(dados_switches)} switch(es)")
    c_sw_count.font = font_total
    c_sw_count.alignment = align_left
    c_sw_count.fill = fill_total
    c_sw_count.border = border_total

    # Coluna 3 (Stack total de membros)
    total_stack_switches = sum(sw.get("stack_count", 1) for sw in dados_switches)
    c_tot_stack = ws.cell(row=current_row, column=3, value=total_stack_switches)
    c_tot_stack.font = font_total
    c_tot_stack.alignment = align_center
    c_tot_stack.fill = fill_total
    c_tot_stack.border = border_total

    # Coluna 4 (Soma de Portas)
    c_tot_portas = ws.cell(row=current_row, column=4, value=total_portas_geral)
    c_tot_portas.font = font_total
    c_tot_portas.alignment = align_center
    c_tot_portas.fill = fill_total
    c_tot_portas.border = border_total

    # Coluna 5 (Soma de APs)
    c_tot_aps = ws.cell(row=current_row, column=5, value=total_aps_geral)
    c_tot_aps.font = font_total
    c_tot_aps.alignment = align_center
    c_tot_aps.fill = fill_total
    c_tot_aps.border = border_total

    # Colunas 6 e 7
    for c_idx in (6, 7):
        c_empty = ws.cell(row=current_row, column=c_idx, value="")
        c_empty.fill = fill_total
        c_empty.border = border_total

    # Auto-ajuste de largura de colunas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Pula título mesclado para não esticar demais
            if cell.row in (1, 2):
                continue
            v_str = str(cell.value or "")
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Ajuste manual fino nas colunas com cabeçalhos longos
    ws.column_dimensions["C"].width = 24  # Total de modulos
    ws.column_dimensions["D"].width = 22  # Oferta total de portas
    ws.column_dimensions["E"].width = 24  # Total de APs

    # 3. Segunda Aba: Enlaces e Conexões Mapeadas
    if links:
        ws_links = wb.create_sheet(title="Enlaces e Conexoes")
        ws_links.views.sheetView[0].showGridLines = True

        ws_links.merge_cells("A1:E1")
        c_title_l = ws_links["A1"]
        c_title_l.value = "ENLACES E PORTAS DE CONEXÃO — TOPOLOGIA MAPEADA"
        c_title_l.font = font_titulo
        c_title_l.alignment = align_left
        ws_links.row_dimensions[1].height = 25

        headers_links = [
            "Dispositivo Local",
            "Porta Local",
            "Dispositivo Vizinho",
            "Porta Remota",
            "Protocolo"
        ]
        ws_links.row_dimensions[3].height = 24
        fill_dark = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        for col_idx, h_text in enumerate(headers_links, 1):
            cell = ws_links.cell(row=3, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_dark
            cell.alignment = align_center

        row_l = 4
        for lk in links:
            ws_links.row_dimensions[row_l].height = 20
            bg_l = fill_alt if row_l % 2 == 1 else None

            c1 = ws_links.cell(row=row_l, column=1, value=lk.get("source_id", ""))
            c2 = ws_links.cell(row=row_l, column=2, value=lk.get("source_port", ""))
            c3 = ws_links.cell(row=row_l, column=3, value=lk.get("target_id", ""))
            c4 = ws_links.cell(row=row_l, column=4, value=lk.get("target_port", ""))
            c5 = ws_links.cell(row=row_l, column=5, value=lk.get("media_type", "LLDP"))

            for c in (c1, c2, c3, c4, c5):
                c.font = font_data
                c.border = border_cell
                if bg_l:
                    c.fill = bg_l

            c1.alignment = align_left
            c2.alignment = align_center
            c3.alignment = align_left
            c4.alignment = align_center
            c5.alignment = align_center

            row_l += 1

        for col in ws_links.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row in (1, 2):
                    continue
                v_str = str(cell.value or "")
                if len(v_str) > max_len:
                    max_len = len(v_str)
            ws_links.column_dimensions[col_letter].width = max(max_len + 4, 18)

    wb.save(caminho_final)
    return caminho_final
