# -*- coding: utf-8 -*-
"""
Módulo de Exportação da Knowledge Base para Word (.docx) e PDF (.pdf).
Gera documentos executivos contendo logotipo corporativo, cabeçalho,
metadados (autor, data), tabelas, blocos de código, alertas, imagens
incorporadas (prints e uploads) e formatação tipográfica profissional.
"""

import os
import re
import html
import base64
from pathlib import Path
from typing import Dict, Any, Optional, List

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.shared import OxmlElement, qn
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    Image as RLImage, HRFlowable, Preformatted
)
from PIL import Image as PILImage

from config import branding
from modules import storage

BASE_DIR = Path(__file__).resolve().parent.parent.parent
DOWNLOADS_DIR = BASE_DIR / "core" / "static" / "downloads"
BRANDING_DIR = BASE_DIR / "core" / "static" / "branding"


def get_company_logo_path() -> Optional[str]:
    """
    Retorna o caminho de um arquivo de logotipo válido (PNG/JPG/JPEG).
    Prioriza o arquivo configurado em branding.LOGO_ARQUIVO.
    Se estiver configurado como SVG, busca o PNG equivalente ou company_logo.png.
    """
    if branding.LOGO_ARQUIVO:
        arq = branding.LOGO_ARQUIVO.strip()
        caminho = BRANDING_DIR / arq
        if caminho.is_file():
            if caminho.suffix.lower() in [".png", ".jpg", ".jpeg"]:
                return str(caminho)
            # Se for SVG, busca PNG correspondente na mesma pasta
            png_equiv = caminho.with_suffix(".png")
            if png_equiv.is_file():
                return str(png_equiv)

    # Fallback para company_logo.png ou company_logo.jpg
    for ext in [".png", ".jpg", ".jpeg"]:
        padrao = BRANDING_DIR / f"company_logo{ext}"
        if padrao.is_file():
            return str(padrao)

    # Busca qualquer imagem raster na pasta branding
    if BRANDING_DIR.is_dir():
        for f in BRANDING_DIR.glob("*.*"):
            if f.suffix.lower() in [".png", ".jpg", ".jpeg"] and f.is_file():
                return str(f)

    return None


def _resolver_caminho_imagem(src: str) -> Optional[str]:
    """Localiza ou decodifica o arquivo da imagem para inserção em Word/PDF."""
    if not src:
        return None

    # Caso 1: URL interna /kb/image/<nome_imagem>
    m_local = re.search(r'/kb/image/([^?#\s"]+)', src)
    if m_local:
        nome_arq = os.path.basename(m_local.group(1))
        caminho = os.path.join(storage.KB_DIR, nome_arq)
        if os.path.isfile(caminho):
            return caminho

    # Caso 2: Arquivo direto dentro de storage.KB_DIR
    caminho_kb = os.path.join(storage.KB_DIR, os.path.basename(src))
    if os.path.isfile(caminho_kb):
        return caminho_kb

    # Caso 3: Imagem Base64 embutida
    if src.startswith("data:image/"):
        try:
            cabecalho, dados_b64 = src.split(",", 1)
            ext_match = re.search(r"data:image/([a-zA-Z0-9]+);base64", cabecalho)
            ext = ext_match.group(1).lower() if ext_match else "png"
            if ext == "jpeg":
                ext = "jpg"
            img_bytes = base64.b64decode(dados_b64)
            temp_name = f"cached_b64_{abs(hash(dados_b64)) % 10000000}.{ext}"
            temp_path = os.path.join(storage.KB_DIR, temp_name)
            if not os.path.isfile(temp_path):
                with open(temp_path, "wb") as f_out:
                    f_out.write(img_bytes)
            return temp_path
        except Exception:
            return None

    return None


# ==============================================================================
# EXPORTAÇÃO MICROSOFT WORD (.DOCX)
# ==============================================================================

def exportar_topico_docx(topico: Dict[str, Any]) -> str:
    """Gera um arquivo .docx executivo completo com logotipo e formatação corporativa."""
    DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Configuração de Margens (2,54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Logotipo Corporativo e Identidade Visual (Tabela de Cabeçalho)
    logo_path = get_company_logo_path()
    empresa_nome = branding.NOME_EMPRESA.strip() or "EMPRESA"

    if logo_path and os.path.isfile(logo_path):
        try:
            head_tab = doc.add_table(rows=1, cols=2)
            head_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            head_tab.autofit = False

            # Coluna Esquerda: Imagem da Logo
            c_logo = head_tab.cell(0, 0)
            c_logo.width = Inches(2.2)
            p_logo = c_logo.paragraphs[0]
            p_logo.paragraph_format.space_after = Pt(2)
            p_logo.add_run().add_picture(logo_path, width=Inches(1.8))

            # Coluna Direita: Nome do Sistema e Empresa
            c_text = head_tab.cell(0, 1)
            c_text.width = Inches(4.3)
            p_text = c_text.paragraphs[0]
            p_text.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_text.paragraph_format.space_after = Pt(2)

            r1 = p_text.add_run(f"{branding.NOME_SISTEMA}\n")
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(15, 23, 42)

            r2 = p_text.add_run(f"{empresa_nome} — Documentação Técnica Oficial")
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = RGBColor(100, 116, 139)
        except Exception:
            p_brand = doc.add_paragraph()
            r_b = p_brand.add_run(f"{branding.NOME_SISTEMA} | {empresa_nome}")
            r_b.font.size = Pt(9)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(100, 116, 139)
    else:
        p_brand = doc.add_paragraph()
        r_b = p_brand.add_run(f"{branding.NOME_SISTEMA} | {empresa_nome}")
        r_b.font.size = Pt(9)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(100, 116, 139)

    # 2. Título do Artigo
    p_title = doc.add_heading(level=1)
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(topico.get("titulo", "Sem Título"))
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # 3. Metadados (Autor e Data)
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(10)
    r_meta = p_meta.add_run(f"Publicado por: {topico.get('autor', 'Admin')}   |   Data: {topico.get('data', 'N/D')}")
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(71, 85, 105)

    # Linha divisória horizontal azul
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    p_div_pPr = p_div._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '0284C7')
    pBdr.append(bottom)
    p_div_pPr.append(pBdr)

    # 4. Processamento do Conteúdo
    raw_md = topico.get("conteudo_md", "")
    lines = raw_md.splitlines()

    in_code_block = False
    code_lines: List[str] = []
    in_table = False
    table_lines: List[str] = []

    def _flush_table(t_lines: List[str]):
        if not t_lines:
            return
        parsed_rows = []
        for l in t_lines:
            # Ignora linha separadora |---|---|
            if re.match(r"^\s*\|?\s*[-:]+[-| :]*\|?\s*$", l):
                continue
            cols = [c.strip() for c in l.strip().strip("|").split("|")]
            if any(cols):
                parsed_rows.append(cols)

        if not parsed_rows:
            return

        max_cols = max(len(r) for r in parsed_rows)
        # Padroniza todas as linhas para ter o mesmo número de colunas
        norm_rows = [r + [""] * (max_cols - len(r)) for r in parsed_rows]

        w_tab = doc.add_table(rows=len(norm_rows), cols=max_cols)
        w_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        w_tab.autofit = True

        for row_idx, r_data in enumerate(norm_rows):
            row = w_tab.rows[row_idx]
            for col_idx, cell_value in enumerate(r_data):
                cell = row.cells[col_idx]
                p_c = cell.paragraphs[0]
                p_c.paragraph_format.space_before = Pt(3)
                p_c.paragraph_format.space_after = Pt(3)
                _format_runs_in_paragraph(p_c, cell_value)

                # Cabeçalho com fundo cinza-claro
                if row_idx == 0:
                    shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shading)
                    for run in p_c.runs:
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                else:
                    for run in p_c.runs:
                        run.font.size = Pt(9)

        # Espaçamento após tabela
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    for line in lines:
        stripped = line.strip()

        # Blocos de Código ```
        if stripped.startswith("```"):
            if in_table:
                _flush_table(table_lines)
                table_lines = []
                in_table = False

            if in_code_block:
                code_text = "\n".join(code_lines)
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.2)
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(6)
                r_c = p_code.add_run(code_text)
                r_c.font.name = "Consolas"
                r_c.font.size = Pt(9)
                r_c.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Tabelas Markdown | col1 | col2 |
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            _flush_table(table_lines)
            table_lines = []
            in_table = False

        # Imagens Markdown ![alt](src)
        m_img = re.search(r"!\[(.*?)\]\((.*?)\)", stripped)
        if m_img:
            alt_text = m_img.group(1).strip()
            src_text = m_img.group(2).strip()
            img_local = _resolver_caminho_imagem(src_text)

            if img_local and os.path.isfile(img_local):
                try:
                    p_img = doc.add_paragraph()
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(2)
                    doc.add_picture(img_local, width=Inches(5.8))

                    # Legenda da figura
                    p_cap = doc.add_paragraph()
                    p_cap.paragraph_format.space_after = Pt(10)
                    r_cap = p_cap.add_run(f"Figura: {alt_text or os.path.basename(img_local)}")
                    r_cap.font.size = Pt(8.5)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = RGBColor(100, 116, 139)
                except Exception:
                    pass
            continue

        # Linha em branco
        if not stripped:
            continue

        # Cabeçalhos com ou sem espaço: # TITULO, #TITULO#, ##SUBTITULO
        m_h3 = re.match(r"^###\s*(.*?)(#+)?$", stripped)
        m_h2 = re.match(r"^##\s*(.*?)(#+)?$", stripped)
        m_h1 = re.match(r"^#\s*(.*?)(#+)?$", stripped)

        if m_h3:
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            _format_runs_in_paragraph(h, m_h3.group(1).strip())
            continue
        elif m_h2:
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            _format_runs_in_paragraph(h, m_h2.group(1).strip())
            continue
        elif m_h1:
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            _format_runs_in_paragraph(h, m_h1.group(1).strip())
            continue

        # Alertas / Callouts (> [!NOTE] ou > texto)
        if stripped.startswith(">"):
            callout_text = re.sub(r"^>\s*(\[!.*\]\s*)?", "", stripped).strip()
            p_callout = doc.add_paragraph()
            p_callout.paragraph_format.left_indent = Inches(0.25)
            p_callout.paragraph_format.space_before = Pt(4)
            p_callout.paragraph_format.space_after = Pt(6)
            r_icon = p_callout.add_run("ℹ ")
            r_icon.font.color.rgb = RGBColor(2, 132, 199)
            _format_runs_in_paragraph(p_callout, callout_text)
            for r in p_callout.runs:
                r.font.italic = True
            continue

        # Listas com marcadores
        if stripped.startswith("- ") or stripped.startswith("* ") or stripped.startswith("• "):
            item_text = re.sub(r"^[-*•]\s+", "", stripped)
            p_li = doc.add_paragraph(style="List Bullet")
            p_li.paragraph_format.space_after = Pt(2)
            _format_runs_in_paragraph(p_li, item_text)
            continue

        # Listas numeradas
        m_num = re.match(r"^\d+\.\s*(.*)$", stripped)
        if m_num:
            item_text = m_num.group(1).strip()
            p_num = doc.add_paragraph(style="List Number")
            p_num.paragraph_format.space_after = Pt(2)
            _format_runs_in_paragraph(p_num, item_text)
            continue

        # Parágrafo comum
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _format_runs_in_paragraph(p, line)

    if in_table:
        _flush_table(table_lines)

    safe_title = re.sub(r'[^A-Za-z0-9_.-]', '_', topico.get("titulo", "artigo"))[:35]
    filename = f"KB_{safe_title}.docx"
    output_path = DOWNLOADS_DIR / filename
    doc.save(str(output_path))
    return str(output_path)


def _format_runs_in_paragraph(paragraph, text: str):
    """Interpreta formatações inline: negrito, itálico, sublinhado e cores de fonte."""
    # Tags <font color="...">
    font_color_match = re.search(r'<font color="(#[0-9a-fA-F]{6}|[a-zA-Z]+)">(.*?)</font>', text)
    if font_color_match:
        c_hex = font_color_match.group(1).lstrip("#")
        inner_text = font_color_match.group(2)
        r = paragraph.add_run(inner_text)
        if len(c_hex) == 6:
            r.font.color.rgb = RGBColor(int(c_hex[0:2], 16), int(c_hex[2:4], 16), int(c_hex[4:6], 16))
        return

    # Marcações Markdown e HTML
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|<u>.*?</u>|<mark>.*?</mark>)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith("<u>") and tok.endswith("</u>"):
            r = paragraph.add_run(tok[3:-4])
            r.underline = True
        elif tok.startswith("<mark>") and tok.endswith("</mark>"):
            r = paragraph.add_run(tok[6:-7])
            r.bold = True
            r.font.color.rgb = RGBColor(180, 83, 9)
        else:
            clean_tok = re.sub(r'<.*?>', '', tok)
            paragraph.add_run(clean_tok)


# ==============================================================================
# EXPORTAÇÃO ADOBE PDF (.PDF)
# ==============================================================================

def exportar_topico_pdf(topico: Dict[str, Any]) -> str:
    """Gera um arquivo .pdf executivo completo com logotipo e formatação corporativa."""
    DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
    safe_title = re.sub(r'[^A-Za-z0-9_.-]', '_', topico.get("titulo", "artigo"))[:35]
    filename = f"KB_{safe_title}.pdf"
    output_path = str(DOWNLOADS_DIR / filename)

    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    normal_style = ParagraphStyle(
        "KBNormal",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=4
    )

    title_style = ParagraphStyle(
        "KBTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=10,
        spaceAfter=4
    )

    h2_style = ParagraphStyle(
        "KBH2",
        parent=styles["Heading2"],
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#0284c7"),
        spaceBefore=12,
        spaceAfter=4
    )

    h3_style = ParagraphStyle(
        "KBH3",
        parent=styles["Heading3"],
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#334155"),
        spaceBefore=8,
        spaceAfter=3
    )

    meta_style = ParagraphStyle(
        "KBMeta",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748b")
    )

    brand_header_style = ParagraphStyle(
        "KBBrandHeader",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0f172a")
    )

    callout_style = ParagraphStyle(
        "KBCallout",
        parent=styles["Normal"],
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#0369a1"),
        backColor=colors.HexColor("#f0f9ff"),
        borderColor=colors.HexColor("#bae6fd"),
        borderWidth=1,
        borderPadding=6,
        spaceBefore=6,
        spaceAfter=6
    )

    code_block_style = ParagraphStyle(
        "KBCodeBlock",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
        backColor=colors.HexColor("#f8fafc"),
        borderColor=colors.HexColor("#e2e8f0"),
        borderWidth=1,
        borderPadding=6,
        spaceBefore=6,
        spaceAfter=8
    )

    story = []

    # 1. Logotipo e Cabeçalho Corporativo no Topo da Página
    logo_path = get_company_logo_path()
    empresa_nome = branding.NOME_EMPRESA.strip() or "EMPRESA"

    if logo_path and os.path.isfile(logo_path):
        try:
            with PILImage.open(logo_path) as pil_img:
                w, h = pil_img.size
                aspect = h / w
                target_w = 130
                target_h = target_w * aspect
                if target_h > 38:
                    target_h = 38
                    target_w = target_h / aspect

            logo_flowable = RLImage(logo_path, width=target_w, height=target_h)
            brand_text = f"<b>{branding.NOME_SISTEMA}</b> &mdash; {html.escape(empresa_nome)}<br/><font size='8' color='#64748b'>Documentação Técnica Oficial</font>"
            header_table = Table(
                [[logo_flowable, Paragraph(brand_text, brand_header_style)]],
                colWidths=[target_w + 12, None]
            )
            header_table.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(header_table)
        except Exception:
            story.append(Paragraph(f"<b>{branding.NOME_SISTEMA}</b> &mdash; {html.escape(empresa_nome)}", brand_header_style))
    else:
        story.append(Paragraph(f"<b>{branding.NOME_SISTEMA}</b> &mdash; {html.escape(empresa_nome)}", brand_header_style))

    # 2. Título do Artigo e Metadados
    story.append(Paragraph(html.escape(topico.get("titulo", "Sem Título")), title_style))
    meta_info = f"Publicado por: <b>{html.escape(topico.get('autor', 'Admin'))}</b> &nbsp;|&nbsp; Data: {html.escape(topico.get('data', 'N/D'))}"
    story.append(Paragraph(meta_info, meta_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0284c7"), spaceAfter=14))

    # 3. Conteúdo Markdown
    raw_md = topico.get("conteudo_md", "")
    lines = raw_md.splitlines()

    in_code_block = False
    code_lines: List[str] = []
    in_table = False
    table_lines: List[str] = []

    def _flush_pdf_table(t_lines: List[str]):
        if not t_lines:
            return
        parsed_rows = []
        for l in t_lines:
            if re.match(r"^\s*\|?\s*[-:]+[-| :]*\|?\s*$", l):
                continue
            cols = [c.strip() for c in l.strip().strip("|").split("|")]
            if any(cols):
                parsed_rows.append(cols)

        if not parsed_rows:
            return

        max_cols = max(len(r) for r in parsed_rows)
        norm_rows = [r + [""] * (max_cols - len(r)) for r in parsed_rows]

        table_data = []
        for r_idx, row in enumerate(norm_rows):
            row_cells = []
            for cell_txt in row:
                if r_idx == 0:
                    c_style = ParagraphStyle(
                        f"TH_{r_idx}",
                        parent=normal_style,
                        fontSize=8.5,
                        leading=11,
                        fontName="Helvetica-Bold",
                        textColor=colors.HexColor("#0f172a")
                    )
                else:
                    c_style = ParagraphStyle(
                        f"TD_{r_idx}",
                        parent=normal_style,
                        fontSize=8.5,
                        leading=11,
                        textColor=colors.HexColor("#334155")
                    )
                row_cells.append(Paragraph(_format_pdf_inlines(cell_txt), c_style))
            table_data.append(row_cells)

        col_width = (doc.width) / max_cols
        pdf_tab = Table(table_data, colWidths=[col_width] * max_cols)
        pdf_tab.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(pdf_tab)
        story.append(Spacer(1, 8))

    for line in lines:
        stripped = line.strip()

        # Blocos de Código ```
        if stripped.startswith("```"):
            if in_table:
                _flush_pdf_table(table_lines)
                table_lines = []
                in_table = False

            if in_code_block:
                code_text = html.escape("\n".join(code_lines))
                story.append(Preformatted(code_text, code_block_style))
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Tabelas Markdown | col1 | col2 |
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            _flush_pdf_table(table_lines)
            table_lines = []
            in_table = False

        # Imagens Markdown ![alt](src)
        m_img = re.search(r"!\[(.*?)\]\((.*?)\)", stripped)
        if m_img:
            alt_text = m_img.group(1).strip()
            src_text = m_img.group(2).strip()
            img_local = _resolver_caminho_imagem(src_text)

            if img_local and os.path.isfile(img_local):
                try:
                    with PILImage.open(img_local) as p_img:
                        iw, ih = p_img.size
                        aspect = ih / iw
                        max_w = doc.width
                        target_w = min(max_w, iw)
                        target_h = target_w * aspect
                        # Limita altura máxima por página
                        if target_h > 420:
                            target_h = 420
                            target_w = target_h / aspect

                    story.append(Spacer(1, 4))
                    story.append(RLImage(img_local, width=target_w, height=target_h))
                    legenda = f"<i><font size='8' color='#64748b'>Figura: {html.escape(alt_text or os.path.basename(img_local))}</font></i>"
                    story.append(Paragraph(legenda, meta_style))
                    story.append(Spacer(1, 6))
                except Exception:
                    pass
            continue

        # Linha em branco
        if not stripped:
            story.append(Spacer(1, 4))
            continue

        # Cabeçalhos com ou sem espaço: # TITULO, #TITULO#, ##SUBTITULO
        m_h3 = re.match(r"^###\s*(.*?)(#+)?$", stripped)
        m_h2 = re.match(r"^##\s*(.*?)(#+)?$", stripped)
        m_h1 = re.match(r"^#\s*(.*?)(#+)?$", stripped)

        if m_h3:
            story.append(Paragraph(_format_pdf_inlines(m_h3.group(1).strip()), h3_style))
            continue
        elif m_h2:
            story.append(Paragraph(_format_pdf_inlines(m_h2.group(1).strip()), h2_style))
            continue
        elif m_h1:
            story.append(Paragraph(_format_pdf_inlines(m_h1.group(1).strip()), title_style))
            continue

        # Alertas / Callouts (> [!NOTE] ou > texto)
        if stripped.startswith(">"):
            callout_text = re.sub(r"^>\s*(\[!.*\]\s*)?", "", stripped).strip()
            p_text = f"<b>Aviso:</b> {_format_pdf_inlines(callout_text)}"
            story.append(Paragraph(p_text, callout_style))
            continue

        # Listas com marcadores
        if stripped.startswith("- ") or stripped.startswith("* ") or stripped.startswith("• "):
            item_text = re.sub(r"^[-*•]\s+", "", stripped)
            bullet_html = f"&bull; &nbsp;{_format_pdf_inlines(item_text)}"
            story.append(Paragraph(bullet_html, normal_style))
            continue

        # Listas numeradas
        m_num = re.match(r"^(\d+\.)\s*(.*)$", stripped)
        if m_num:
            num_html = f"<b>{m_num.group(1)}</b> &nbsp;{_format_pdf_inlines(m_num.group(2))}"
            story.append(Paragraph(num_html, normal_style))
            continue

        # Parágrafo comum
        story.append(Paragraph(_format_pdf_inlines(stripped), normal_style))

    if in_table:
        _flush_pdf_table(table_lines)

    doc.build(story)
    return output_path


def _format_pdf_inlines(text: str) -> str:
    """Converte formatação de texto Markdown e tags HTML para o formato suportado pelo ReportLab."""
    # Preserva tags de cor <font color="...">
    color_map: Dict[str, str] = {}
    def _save_font(match):
        cid = f"__FONT_COLOR_{len(color_map)}__"
        color_map[cid] = (match.group(1), match.group(2))
        return cid

    clean_text = re.sub(r'<font color="(#[0-9a-fA-F]{6}|[a-zA-Z]+)">(.*?)</font>', _save_font, text)

    # Escapa caracteres HTML
    escaped = html.escape(clean_text)

    # Converte negrito, itálico, sublinhado
    escaped = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', escaped)
    escaped = re.sub(r'\*(.*?)\*', r'<i>\1</i>', escaped)
    escaped = re.sub(r'&lt;u&gt;(.*?)&lt;/u&gt;', r'<u>\1</u>', escaped)
    escaped = re.sub(r'&lt;mark&gt;(.*?)&lt;/mark&gt;', r'<b><font color="#b45309">\1</font></b>', escaped)

    # Restaura cores de fonte
    for cid, (cor, conteudo) in color_map.items():
        escaped = escaped.replace(cid, f'<font color="{cor}">{html.escape(conteudo)}</font>')

    return escaped
