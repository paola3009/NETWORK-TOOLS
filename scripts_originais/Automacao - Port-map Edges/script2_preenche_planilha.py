
#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import json
import pandas as pd
from typing import Tuple, Optional, List
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PLANILHA = "Port-Map-Base.xlsx"
JSON_FILE = "dados_portmap.json"
OUT_FILE = "Port-Map-Base-Preenchida.xlsx"

DEBUG = False  # True para logs detalhados

# ---------------------------
# Utilitários de normalização
# ---------------------------
def as_str(x) -> str:
    if isinstance(x, str):
        return x
    try:
        s = str(x)
    except Exception:
        return ""
    return "" if s.lower() == "nan" else s

def label(h: Optional[str]) -> str:
    """Rótulo sem domínio, em MAIÚSCULAS (p/ comparação)."""
    h = as_str(h).strip()
    if not h:
        return ""
    return h.split(".")[0].upper()

def is_switch(h: Optional[str]) -> bool:
    return label(h).startswith("SW")

# ---------------------------
# Índices de busca (a partir do JSON)
# ---------------------------
def build_indexes(dados: dict):
    """
    pairs_by_name[(SRC_LABEL, DST_LABEL)] = (local_intf, neighbor_intf)
    pairs_by_ip[(SRC_LABEL, DST_IP)]     = (local_intf, neighbor_intf)
    loopbacks[label(HOST)]               = ip_loopback
    """
    pairs_by_name = {}
    pairs_by_ip = {}
    loopbacks = {}

    for host_key, obj in dados.items():
        src = label(host_key)
        loopbacks[src] = as_str(obj.get("loopback0_ip") or "")

        for n in obj.get("cdp", []):
            neigh_name = as_str(n.get("neighbor"))
            neigh_ip   = as_str(n.get("neighbor_ip") or "")
            li = as_str(n.get("local_intf") or "")
            ri = as_str(n.get("neighbor_intf") or "")
            dst = label(neigh_name)

            if src and dst:
                pairs_by_name[(src, dst)] = (li, ri)
            if src and neigh_ip:
                pairs_by_ip[(src, neigh_ip)] = (li, ri)

    return pairs_by_name, pairs_by_ip, loopbacks

# ---------------------------
# Match helpers
# ---------------------------
def find_by_name(A: str, B: str, pairs_by_name: dict) -> Tuple[Optional[str], Optional[str]]:
    return pairs_by_name.get((label(A), label(B)), (None, None))

def find_by_name_reverse(A: str, B: str, pairs_by_name: dict) -> Tuple[Optional[str], Optional[str]]:
    li, ri = pairs_by_name.get((label(B), label(A)), (None, None))
    if li or ri:
        # invertendo p/ preencher como A,B
        return ri, li
    return None, None

def find_by_ip(A: str, ipB: str, pairs_by_ip: dict) -> Tuple[Optional[str], Optional[str]]:
    return pairs_by_ip.get((label(A), as_str(ipB)), (None, None))

def find_by_ip_reverse(B: str, ipA: str, pairs_by_ip: dict) -> Tuple[Optional[str], Optional[str]]:
    li, ri = pairs_by_ip.get((label(B), as_str(ipA)), (None, None))
    if li or ri:
        return ri, li  # invertendo p/ preencher como A,B
    return None, None

# ---------------------------
# Mesclagem pós-processamento
# ---------------------------
def _header_to_col_index(ws, header_name: str) -> Optional[int]:
    for col in range(1, ws.max_column + 1):
        if ws[f"{get_column_letter(col)}1"].value == header_name:
            return col
    return None

def merge_blocks_by_key(excel_file: str, key_col_name: str, cols_to_merge: List[str]):
    """
    Reabre o Excel e, usando a coluna 'key_col_name' como chave de bloco
    (linhas consecutivas com mesmo valor), mescla as células de:
      - a própria 'key_col_name'
      - todas as colunas listadas em 'cols_to_merge'
    em cada bloco encontrado.
    """
    wb = load_workbook(excel_file)
    ws = wb.active

    key_col_idx = _header_to_col_index(ws, key_col_name)
    if key_col_idx is None:
        print(f"Coluna-chave '{key_col_name}' não encontrada para mesclagem.")
        wb.save(excel_file)
        return

    merge_col_indices = []
    for name in cols_to_merge:
        idx = _header_to_col_index(ws, name)
        if idx is not None:
            merge_col_indices.append(idx)
        else:
            print(f"(i) Coluna '{name}' não encontrada; mesclagem ignorada para ela.")

    start = 2  # dados iniciam na linha 2
    while start <= ws.max_row:
        key_val = ws.cell(row=start, column=key_col_idx).value
        end = start + 1
        while end <= ws.max_row and ws.cell(row=end, column=key_col_idx).value == key_val:
            end += 1

        # se bloco possui >1 linha, mesclar
        if end - start > 1:
            # mescla a própria coluna-chave
            ws.merge_cells(start_row=start, start_column=key_col_idx,
                           end_row=end-1, end_column=key_col_idx)
            # mescla demais colunas-alvo dentro do mesmo bloco
            for c in merge_col_indices:
                ws.merge_cells(start_row=start, start_column=c,
                               end_row=end-1, end_column=c)

        start = end

    wb.save(excel_file)
    print(f"✔ Mesclagem por bloco de '{key_col_name}' concluída para colunas: {[key_col_name] + cols_to_merge}")

# ---------------------------
# Principal
# ---------------------------
def main():
    df = pd.read_excel(PLANILHA)

    # Corrige células mescladas da base em "Hostname B"
    if "Hostname B" in df.columns:
        df["Hostname B"] = df["Hostname B"].ffill()

    # Força colunas de saída como texto
    for col in ["IP Address A", "IP Address B", "Interface A", "Interface B"]:
        if col in df.columns:
            df[col] = df[col].astype("object")

    # Carrega JSON e cria índices
    with open(JSON_FILE, "r", encoding="utf-8") as f:
        dados = json.load(f)
    pairs_by_name, pairs_by_ip, loopbacks = build_indexes(dados)

    not_sw_pairs = []
    not_found = []

    for idx, row in df.iterrows():
        hostA = as_str(row.get("Hostname A"))
        hostB = as_str(row.get("Hostname B"))

        # IPs de loopback para A e B
        df.at[idx, "IP Address A"] = loopbacks.get(label(hostA), "")
        df.at[idx, "IP Address B"] = loopbacks.get(label(hostB), "")

        # Apenas SW<->SW preenche interface
        if not (is_switch(hostA) and is_switch(hostB)):
            not_sw_pairs.append((hostA, hostB))
            continue

        ipA = as_str(df.at[idx, "IP Address A"])
        ipB = as_str(df.at[idx, "IP Address B"])

        # 1) A->B por nome
        la, nb = find_by_name(hostA, hostB, pairs_by_name)

        # 2) B->A por nome (espelhado)
        if not la and not nb:
            la, nb = find_by_name_reverse(hostA, hostB, pairs_by_name)

        # 3) A->B por IP
        if not la and not nb and ipB:
            la, nb = find_by_ip(hostA, ipB, pairs_by_ip)

        # 4) B->A por IP (espelhado)
        if not la and not nb and ipA:
            la, nb = find_by_ip_reverse(hostB, ipA, pairs_by_ip)

        if la or nb:
            if la:
                df.at[idx, "Interface A"] = la
            if nb:
                df.at[idx, "Interface B"] = nb
        else:
            not_found.append((hostA, hostB))

    # Salva a planilha preenchida (sem mesclagens)
    df.to_excel(OUT_FILE, index=False)
    print(f"\n✔ Planilha '{OUT_FILE}' gerada.")

    # ---- Remesclagem por blocos de Hostname B ----
    # Mesclar também:
    # - IP Address B  (coluna G na sua planilha)
    # - Prédio        (J)
    # - Andar         (K)
    # - Bastidor      (L)
    merge_blocks_by_key(
        excel_file=OUT_FILE,
        key_col_name="Hostname B",
        cols_to_merge=["IP Address B", "Prédio", "Andar", "Bastidor"]
    )

    # Resumo
    if not_sw_pairs:
        print("\n(i) Linhas ignoradas (não SW<->SW):")
        for a, b in not_sw_pairs:
            print(f"   - {a} <-> {b}")

    if not_found:
        print("\n(⚠) Pares SW<->SW não mapeados após as tentativas:")
        for a, b in not_found:
            print(f"   - {a} <-> {b}")

if __name__ == "__main__":
    main()
